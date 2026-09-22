"""Builds sample_data.xlsx (the Excel template). Run: python make_template.py"""
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.comments import Comment

FONT = "Arial"
HEAD_FILL = PatternFill("solid", start_color="1F3864")
HELP_FILL = PatternFill("solid", start_color="E7E6E6")
thin = Side(style="thin", color="BFBFBF")
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)
LAST = 500  # rows covered by dropdowns

POS_HEADERS = ["position_id", "position_title", "department", "reports_to",
               "status", "email", "phone", "sort_order", "reports_to_title (auto - do not type)"]
POS_WIDTHS = [12, 32, 20, 12, 10, 26, 16, 11, 34]

positions = [
    ("P001", "Director", "Administration", "", "Filled", "director@example.org", "9800000001", 1),
    ("P002", "Deputy Director", "Administration", "P001", "Filled", "", "", 1),
    ("P003", "Finance Manager", "Finance", "P001", "Filled", "finance@example.org", "9800000003", 2),
    ("P004", "HR Manager", "HR", "P001", "Vacant", "", "", 3),
    ("P005", "IT Manager", "IT", "P001", "Filled", "", "", 4),
    ("P006", "Administrative Officer", "Administration", "P002", "Filled", "", "", 1),
    ("P007", "Accountant", "Finance", "P003", "Filled", "", "", 1),
    ("P008", "Audit Officer", "Finance", "P003", "Filled", "", "", 2),
    ("P009", "Cashier", "Finance", "P007", "Filled", "", "", 1),
    ("P010", "HR Assistant", "HR", "P004", "Filled", "", "", 1),
    ("P011", "System Administrator", "IT", "P005", "Filled", "", "", 1),
    ("P012", "Web Developer", "IT", "P005", "Vacant", "", "", 2),
    ("P013", "Office Assistant", "Administration", "P006", "Filled", "", "", 1),
]
extras = [
    ("P007", "P002", "Dotted line"),
    ("P010", "P003", "Functional"),
    ("P011", "P002", "Project"),
]


def style_header(ws, row, ncols):
    for c in range(1, ncols + 1):
        cell = ws.cell(row=row, column=c)
        cell.font = Font(name=FONT, bold=True, color="FFFFFF")
        cell.fill = HEAD_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = BORDER


wb = Workbook()

# ---------------- Instructions ----------------
ws = wb.active
ws.title = "Instructions"
ws.column_dimensions["A"].width = 26
ws.column_dimensions["B"].width = 95
lines = [
    ("ORGANOGRAM DATA TEMPLATE", None),
    ("", None),
    ("How to use", "Fill the 'Positions' sheet (one row per position) and, only if needed, the 'Extra_Reporting' sheet. Then upload this file in the app."),
    ("Sample rows", "Rows already in the two sheets are SAMPLE DATA for a made-up institute. Delete or overwrite them with your own."),
    ("Cells to edit", "Positions: columns A to H.  Extra_Reporting: columns A to C.  Column I of Positions is filled automatically - do not type there."),
    ("", None),
    ("POSITIONS SHEET", None),
    ("position_id", "REQUIRED. Unique code such as P001, P002. Never repeat, reuse or change it once other rows refer to it."),
    ("position_title", "REQUIRED. Name of the post, e.g. Finance Manager."),
    ("department", "Used for the box colours in the chart."),
    ("reports_to", "position_id of the MAIN (solid line) manager. Leave BLANK for top-level (Level 1) posts. You may have many Level 1 posts."),
    ("status", "Filled or Vacant (choose from the dropdown). Vacant posts are shown with a red dashed border."),
    ("email, phone", "Optional. Shown in the Word position table and, if you tick the option, in the chart boxes."),
    ("sort_order", "Optional number. Controls left-to-right order among posts under the same manager."),
    ("Level", "Do NOT enter it. The app works it out from the reporting chain."),
    ("", None),
    ("EXTRA_REPORTING SHEET", None),
    ("position_id", "The post that has a second (or third) manager."),
    ("also_reports_to", "position_id of the additional manager. Add one row per extra manager."),
    ("relationship", "Optional: Dotted line, Functional, Acting or Project. Drawn as a dashed line."),
    ("", None),
    ("RULES THE APP CHECKS", None),
    ("1", "Every position_id is unique and not blank."),
    ("2", "Every reports_to / also_reports_to value exists in the Positions sheet."),
    ("3", "No post reports to itself and there are no loops (A reports to B, B reports to A)."),
    ("4", "At least one top-level post exists."),
    ("5", "The same manager is not listed as both main and extra manager for one post."),
]
for r, (a, b) in enumerate(lines, start=1):
    ca = ws.cell(row=r, column=1, value=a)
    ca.font = Font(name=FONT, bold=True)
    if b is not None:
        cb = ws.cell(row=r, column=2, value=b)
        cb.font = Font(name=FONT)
        cb.alignment = Alignment(wrap_text=True, vertical="top")
        ca.alignment = Alignment(vertical="top")
    elif a:
        ca.fill = HELP_FILL
        ws.cell(row=r, column=2).fill = HELP_FILL
ws["A1"].font = Font(name=FONT, bold=True, size=14)
ws["A1"].fill = PatternFill(fill_type=None)

# ---------------- Positions ----------------
wp = wb.create_sheet("Positions")
for c, (h, w) in enumerate(zip(POS_HEADERS, POS_WIDTHS), start=1):
    wp.cell(row=1, column=c, value=h)
    wp.column_dimensions[chr(64 + c)].width = w
style_header(wp, 1, len(POS_HEADERS))
wp.row_dimensions[1].height = 32
wp["I1"].fill = PatternFill("solid", start_color="7F7F7F")

for r, row in enumerate(positions, start=2):
    for c, v in enumerate(row, start=1):
        wp.cell(row=r, column=c, value=(v if v != "" else None))
for r in range(2, 301):
    for c in range(1, 9):
        cell = wp.cell(row=r, column=c)
        cell.font = Font(name=FONT)
        cell.border = BORDER
    j = wp.cell(row=r, column=9,
                value=f'=IF(D{r}="","",IFERROR(INDEX($B$2:$B${LAST},MATCH(D{r},$A$2:$A${LAST},0)),"NOT FOUND"))')
    j.font = Font(name=FONT, color="7F7F7F", italic=True)
    j.fill = HELP_FILL
    j.border = BORDER
wp.freeze_panes = "C2"

dv_status = DataValidation(type="list", formula1='"Filled,Vacant"', allow_blank=True)
dv_status.error = "Choose Filled or Vacant"
dv_status.errorTitle = "Invalid status"
wp.add_data_validation(dv_status)
dv_status.add(f"E2:E{LAST}")

dv_parent = DataValidation(type="list", formula1=f"=$A$2:$A${LAST}", allow_blank=True)
dv_parent.error = "Choose an existing position_id from column A (leave blank for a top-level post)."
dv_parent.errorTitle = "Unknown position_id"
wp.add_data_validation(dv_parent)
dv_parent.add(f"D2:D{LAST}")

dv_order = DataValidation(type="whole", operator="greaterThanOrEqual", formula1="0", allow_blank=True)
dv_order.error = "sort_order must be a whole number"
wp.add_data_validation(dv_order)
dv_order.add(f"H2:H{LAST}")

wp["D1"].comment = Comment("Leave blank for top-level (Level 1) posts. Otherwise pick the position_id of the main manager.", "Template")

# ---------------- Extra_Reporting ----------------
we = wb.create_sheet("Extra_Reporting")
for c, (h, w) in enumerate(zip(["position_id", "also_reports_to", "relationship"], [14, 18, 20]), start=1):
    we.cell(row=1, column=c, value=h)
    we.column_dimensions[chr(64 + c)].width = w
style_header(we, 1, 3)
for r, row in enumerate(extras, start=2):
    for c, v in enumerate(row, start=1):
        we.cell(row=r, column=c, value=v)
for r in range(2, 201):
    for c in range(1, 4):
        cell = we.cell(row=r, column=c)
        cell.font = Font(name=FONT)
        cell.border = BORDER
we.freeze_panes = "A2"

dv_a = DataValidation(type="list", formula1=f"=Positions!$A$2:$A${LAST}", allow_blank=True)
dv_a.error = "Choose an existing position_id from the Positions sheet."
we.add_data_validation(dv_a)
dv_a.add(f"A2:B{LAST}")
dv_rel = DataValidation(type="list", formula1='"Dotted line,Functional,Acting,Project"', allow_blank=True)
we.add_data_validation(dv_rel)
dv_rel.add(f"C2:C{LAST}")

wb.save("sample_data.xlsx")
print("saved")
