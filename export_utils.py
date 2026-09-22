"""PDF and Word export."""
from __future__ import annotations

import datetime
import io
from typing import Dict, List

import pandas as pd
from docx import Document
from docx.enum.section import WD_ORIENT, WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from PIL import Image
from pypdf import PageObject, PdfReader, PdfWriter, Transformation

from data_utils import compute_levels, positions_by_id, tree_rows

# paper sizes in millimetres, landscape (width, height)
PAPERS = {"A4": (297, 210), "A3": (420, 297), "A2": (594, 420)}
MM_TO_PT = 72 / 25.4


# ----------------------------------------------------------------------------
# PDF
# ----------------------------------------------------------------------------
def _is_portrait(width: float, height: float) -> bool:
    """Tall charts (e.g. left-to-right layouts) go on portrait sheets."""
    return height > width * 1.15


def pages_to_pdf(pages: List[dict], paper: str = "A3", margin_mm: float = 10) -> bytes:
    """Every chart page becomes one sheet of the chosen paper size (landscape, or portrait
    for tall charts). The chart stays vector (sharp at any zoom); it is scaled to fit."""
    m = margin_mm * MM_TO_PT
    writer = PdfWriter()
    for p in pages:
        src = PdfReader(io.BytesIO(p["graph"].pipe(format="pdf"))).pages[0]
        sw, sh = float(src.mediabox.width), float(src.mediabox.height)
        long_side, short_side = (v * MM_TO_PT for v in PAPERS[paper])
        pw, ph = (short_side, long_side) if _is_portrait(sw, sh) else (long_side, short_side)
        scale = min((pw - 2 * m) / sw, (ph - 2 * m) / sh, 1.6)  # don't blow tiny charts up too much
        sheet = PageObject.create_blank_page(width=pw, height=ph)
        tx = (pw - sw * scale) / 2
        ty = (ph - sh * scale) / 2
        sheet.merge_transformed_page(src, Transformation().scale(scale).translate(tx, ty))
        writer.add_page(sheet)
    out = io.BytesIO()
    writer.write(out)
    return out.getvalue()


def graph_to_png(graph, max_pixels: int = 7000, dpi: int = 200) -> bytes:
    """PNG of a graph. Resolution is lowered for very wide charts so the picture stays a sane size."""
    pdf = PdfReader(io.BytesIO(graph.pipe(format="pdf"))).pages[0]
    width_in = float(pdf.mediabox.width) / 72
    height_in = float(pdf.mediabox.height) / 72
    use_dpi = max(60, min(dpi, int(max_pixels / max(width_in, height_in, 1))))
    g = graph.copy()
    g.graph_attr["dpi"] = str(use_dpi)
    return g.pipe(format="png")


# ----------------------------------------------------------------------------
# Word
# ----------------------------------------------------------------------------
def _shade(cell, hex_fill: str):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_fill)
    tcPr.append(shd)


def _repeat_header(row):
    trPr = row._tr.get_or_add_trPr()
    el = OxmlElement("w:tblHeader")
    el.set(qn("w:val"), "true")
    trPr.append(el)


def _no_split(row):
    trPr = row._tr.get_or_add_trPr()
    el = OxmlElement("w:cantSplit")
    el.set(qn("w:val"), "true")
    trPr.append(el)


def _set_cell(cell, text, bold=False, size=8, color=None):
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    run = p.add_run(text)
    run.font.size = Pt(size)
    run.bold = bold
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def _set_section(sec, paper: str, portrait: bool):
    long_mm, short_mm = PAPERS[paper]
    w, h = (short_mm, long_mm) if portrait else (long_mm, short_mm)
    sec.orientation = WD_ORIENT.PORTRAIT if portrait else WD_ORIENT.LANDSCAPE
    sec.page_width, sec.page_height = Cm(w / 10), Cm(h / 10)
    for side in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(sec, side, Cm(1.5))
    return w, h


def build_word(pages: List[dict], pos: pd.DataFrame, extra: pd.DataFrame,
               paper: str = "A3", institution: str = "", include_table: bool = True) -> bytes:
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Arial"
    style.font.size = Pt(10)

    for i, p in enumerate(pages):
        png = graph_to_png(p["graph"])
        wpx, hpx = Image.open(io.BytesIO(png)).size
        portrait = _is_portrait(wpx, hpx)
        if i == 0:
            sec = doc.sections[0]
        else:
            sec = doc.add_section(WD_SECTION.NEW_PAGE)
        w_mm, h_mm = _set_section(sec, paper, portrait)
        avail_w = Cm(w_mm / 10 - 3.0)

        if i == 0 and institution:
            t = doc.add_paragraph()
            r = t.add_run(institution)
            r.bold, r.font.size = True, Pt(20)
            t.alignment = WD_ALIGN_PARAGRAPH.CENTER
            t.paragraph_format.space_after = Pt(0)
        if i == 0:
            sub = doc.add_paragraph()
            sr = sub.add_run("Organisational chart - " + datetime.date.today().strftime("%d %B %Y"))
            sr.font.size = Pt(10)
            sr.font.color.rgb = RGBColor(0x59, 0x59, 0x59)
            sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
        if len(pages) > 1 or p["title"] != "Organogram":
            h = doc.add_paragraph()
            hr = h.add_run(f"Page {i + 1}: {p['title']}")
            hr.bold, hr.font.size = True, Pt(12)

        avail_h = Cm(h_mm / 10 - 3.0 - 3.6)  # margins + heading lines
        scale = min(avail_w / wpx, avail_h / hpx)
        pic_par = doc.add_paragraph()
        pic_par.alignment = WD_ALIGN_PARAGRAPH.CENTER
        pic_par.add_run().add_picture(io.BytesIO(png), width=int(wpx * scale), height=int(hpx * scale))

    if include_table and not pos.empty:
        _set_section(doc.add_section(WD_SECTION.NEW_PAGE), paper, portrait=False)
        head = doc.add_paragraph()
        hr = head.add_run("Position details")
        hr.bold, hr.font.size = True, Pt(14)

        by_id = positions_by_id(pos)
        levels = compute_levels(pos)
        extras: Dict[str, List[str]] = {}
        for r in extra.to_dict("records"):
            m = by_id.get(r["also_reports_to"])
            label = r["also_reports_to"] + (f" {m['position_title']}" if m else "")
            if r["relationship"]:
                label += f" [{r['relationship']}]"
            extras.setdefault(r["position_id"], []).append(label)

        has_names = (pos["name"] != "").any()
        cols = ["Level", "ID", "Position"] + (["Name"] if has_names else []) + [
            "Department", "Status", "Reports to", "Also reports to", "Email", "Phone"]
        widths_cm = [1.5, 1.4, 5.5] + ([4.0] if has_names else []) + [3.4, 1.8, 4.6, 5.0, 4.5, 2.8]
        table = doc.add_table(rows=1, cols=len(cols))
        table.style = "Table Grid"
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.autofit = False
        for c, name in enumerate(cols):
            _set_cell(table.rows[0].cells[c], name, bold=True, size=8, color="FFFFFF")
            _shade(table.rows[0].cells[c], "1F3864")
        _repeat_header(table.rows[0])

        for pid, depth in tree_rows(pos):
            row = by_id[pid]
            parent = by_id.get(row["reports_to"])
            cells = table.add_row().cells
            status = row["status"] or "Filled"
            vals = [str(levels.get(pid, "")), pid, ("    " * depth) + row["position_title"]]
            if has_names:
                vals.append(row["name"] or "-")
            vals += [
                row["department"], status,
                (f"{parent['position_id']} {parent['position_title']}" if parent else "-"),
                "; ".join(extras.get(pid, [])) or "-", row["email"], row["phone"],
            ]
            status_col = cols.index("Status")
            for c, v in enumerate(vals):
                _set_cell(cells[c], v, size=8, color="C00000" if (c == status_col and v == "Vacant") else None)
            _no_split(table.rows[-1])

        # fit the table to the page width, then set both the column grid and each cell
        long_mm = PAPERS[paper][0]
        factor = (long_mm / 10 - 3.0) / sum(widths_cm)
        widths = [Cm(w * factor) for w in widths_cm]
        for c, w in enumerate(widths):
            table.columns[c].width = w
        for row in table.rows:
            for c, w in enumerate(widths):
                row.cells[c].width = w

    out = io.BytesIO()
    doc.save(out)
    return out.getvalue()
